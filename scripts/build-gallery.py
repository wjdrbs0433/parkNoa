"""Generate Gallery page data and small WebP previews from the Excel workbook."""

import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image, ImageOps


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'content' / 'site-content.xlsx'
IMAGE_DIR = ROOT / 'img'
THUMB_DIR = IMAGE_DIR / 'gallery-thumbs'
PREVIEW_DIR = IMAGE_DIR / 'gallery-previews'
OUTPUT = ROOT / 'inc' / 'content' / 'gallery-manifest.js'
HEADERS = ('파일명', '파일형식', '년도', '월', '순서', '그리드 표시', '사진 설명 (선택)', '캐러셀 순서 (1~3)')
EXTENSIONS = {'jpg', 'jpeg', 'png', 'webp'}
MONTHS = ('Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec')


def as_text(value):
    return '' if value is None else str(value).strip()


def as_integer(value, label, row_number, minimum, maximum):
    if isinstance(value, bool) or not isinstance(value, (int, float)) or int(value) != value:
        raise ValueError(f'Gallery {row_number}행: {label}은 숫자로 입력하세요.')
    number = int(value)
    if not minimum <= number <= maximum:
        raise ValueError(f'Gallery {row_number}행: {label}은 {minimum}~{maximum} 범위여야 합니다.')
    return number


def as_flag(value, label, row_number, default):
    value = as_text(value).lower()
    if not value:
        return default
    if value in ('예', 'y', 'yes', 'true', '1'):
        return True
    if value in ('아니요', 'n', 'no', 'false', '0'):
        return False
    raise ValueError(f'Gallery {row_number}행: {label}은 예 또는 아니요로 입력하세요.')


def build_data(workbook, image_dir=IMAGE_DIR):
    if 'Gallery' not in workbook:
        raise ValueError('Gallery 시트가 없습니다.')
    rows = workbook['Gallery'].iter_rows(values_only=True)
    if tuple(next(rows)[:8]) != HEADERS:
        raise ValueError('Gallery 시트의 A~H열 제목이 변경되었습니다.')

    items = []
    skipped = []
    filenames = set()
    positions = set()
    carousel_orders = set()
    for row_number, row in enumerate(rows, start=2):
        values = list(row[:8]) + [None] * max(0, 8 - len(row))
        if not any(as_text(value) for value in values):
            continue
        stem, extension = map(as_text, values[:2])
        extension = extension.lower().lstrip('.')
        if not stem or stem in ('.', '..') or '/' in stem or '\\' in stem or not extension or '.' in stem:
            raise ValueError(f'Gallery {row_number}행: 파일명에는 확장자와 경로를 빼고 입력하세요.')
        if extension not in EXTENSIONS:
            raise ValueError(f'Gallery {row_number}행: 지원하지 않는 파일형식 {extension!r}')
        filename = f'{stem}.{extension}'
        if not (image_dir / filename).is_file():
            alternatives = sorted(path.name for path in image_dir.glob(f'{stem}.*') if path.is_file())
            hint = f' 실제 파일: {", ".join(alternatives)}.' if alternatives else ''
            skipped.append(
                f'Gallery {row_number}행: img/{filename} 파일이 없습니다.'
                f'{hint} 이 사진은 페이지에서 제외합니다.'
            )
            continue
        if filename.lower() in filenames:
            raise ValueError(f'Gallery {row_number}행: 중복 파일명 {filename}')
        filenames.add(filename.lower())
        year = as_integer(values[2], '년도', row_number, 1000, 9999)
        month = as_integer(values[3], '월', row_number, 1, 12)
        order = as_integer(values[4], '순서', row_number, 1, 9999)
        position = (year, month, order)
        if position in positions:
            raise ValueError(f'Gallery {row_number}행: 같은 연월의 순서 {order}가 중복됐습니다.')
        positions.add(position)
        visible = as_flag(values[5], '그리드 표시', row_number, True)
        description = as_text(values[6]) or f'{MONTHS[month - 1]} {year}'
        carousel_order = None
        if as_text(values[7]):
            carousel_order = as_integer(values[7], '캐러셀 순서', row_number, 1, 3)
            if carousel_order in carousel_orders:
                raise ValueError(f'Gallery {row_number}행: 캐러셀 순서 {carousel_order}가 중복됐습니다.')
            carousel_orders.add(carousel_order)
        items.append({
            'year': year, 'month': month, 'order': order,
            'src': f'img/{filename}', 'label': description,
            'visible': visible, 'carouselOrder': carousel_order,
        })

    items.sort(key=lambda item: (item['year'], item['month'], item['order']), reverse=True)
    return {'items': items, 'skipped': skipped}


def file_digest(path):
    digest = hashlib.sha256()
    with path.open('rb') as image_file:
        for chunk in iter(lambda: image_file.read(1024 * 1024), b''):
            digest.update(chunk)
    return digest.hexdigest()[:10]


def save_webp(source, destination, width, quality):
    if destination.exists():
        return
    with Image.open(source) as original:
        image = ImageOps.exif_transpose(original)
        image.thumbnail((width, width), Image.Resampling.LANCZOS)
        image = image.convert('RGBA' if 'A' in image.getbands() else 'RGB')
        temporary = destination.with_name(destination.stem + '.tmp.webp')
        image.save(temporary, format='WEBP', quality=quality, method=4)
        temporary.replace(destination)


def main():
    workbook = load_workbook(SOURCE, read_only=True, data_only=True)
    try:
        data = build_data(workbook)
    finally:
        workbook.close()
    skipped = data.pop('skipped')
    for warning in skipped:
        print(warning)

    THUMB_DIR.mkdir(exist_ok=True)
    PREVIEW_DIR.mkdir(exist_ok=True)
    for item in data['items']:
        if not (item['visible'] or item['carouselOrder'] is not None):
            continue
        source = ROOT / item['src']
        stem = source.stem
        digest = file_digest(source)
        filename = f'{stem}-{digest}.webp'
        thumb = THUMB_DIR / filename
        preview = PREVIEW_DIR / filename
        save_webp(source, thumb, 640, 76)
        save_webp(source, preview, 1200, 78)
        item['thumb'] = f'img/gallery-thumbs/{filename}'
        item['preview'] = f'img/gallery-previews/{filename}'

    text = '// Generated from content/site-content.xlsx by scripts/build-gallery.py.\n'
    text += 'window.GALLERY_CONTENT = ' + json.dumps(data, ensure_ascii=False, indent=2) + ';\n'
    temporary = OUTPUT.with_name(OUTPUT.name + '.tmp')
    temporary.write_text(text, encoding='utf-8')
    temporary.replace(OUTPUT)
    visible = sum(item['visible'] for item in data['items'])
    print(f'Generated {OUTPUT.relative_to(ROOT)} from {len(data["items"])} photos ({visible} in the grid)')


if __name__ == '__main__':
    main()
