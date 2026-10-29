"""Generate the About page data from content/site-content.xlsx."""

import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'content' / 'site-content.xlsx'
OUTPUT = ROOT / 'inc' / 'content' / 'about-content.js'


FIELDS = {
    '프로필 사진': ('profilePhoto', 'src'),
    '사진 설명': ('profilePhoto', 'alt'),
    '이름': ('hero', 'name'),
    '소속·직책': ('hero', 'position'),
    '부서': ('hero', 'metaDepartment'),
    '이전 학교': ('hero', 'metaPriorSchool'),
    '연구 소개': ('hero', 'highlight'),
    '소개 제목': ('biography', 'sectionTitle'),
    'Google Scholar 주소': ('biography', 'scholarLink', 'url'),
    'Google Scholar 문구': ('biography', 'scholarLink', 'label'),
    '연구 분야 제목': ('cv', 'researchInterests', 'title'),
    '학력 제목': ('cv', 'education', 'title'),
    '심사 저널 제목': ('cv', 'peerReview', 'title'),
    '심사 기간': ('cv', 'peerReview', 'period'),
    '심사 역할': ('cv', 'peerReview', 'roleLabel'),
    '활동 제목': ('cv', 'leadership', 'title'),
}

CATEGORIES = {'소개 문단', '연구 분야', '학력', '심사 저널', '활동'}


def as_text(value):
    return '' if value is None else str(value).strip()


def put_nested(data, path, value):
    current = data
    for part in path[:-1]:
        current = current.setdefault(part, {})
    current[path[-1]] = value


def build_data(workbook):
    if 'About' not in workbook or 'About 목록' not in workbook:
        raise ValueError('About 또는 About 목록 시트가 없습니다.')

    settings = workbook['About'].iter_rows(values_only=True)
    if tuple(next(settings)[:2]) != ('항목', '입력 내용'):
        raise ValueError('About 시트의 열 제목이 변경되었습니다.')

    data = {}
    seen = set()
    for row_number, row in enumerate(settings, start=2):
        label = as_text(row[0])
        if not label:
            continue
        if label not in FIELDS:
            raise ValueError(f'About 시트 {row_number}행: 알 수 없는 항목 {label!r}')
        if label in seen:
            raise ValueError(f'About 시트 {row_number}행: 중복 항목 {label!r}')
        seen.add(label)
        put_nested(data, FIELDS[label], as_text(row[1]))

    missing = set(FIELDS) - seen
    if missing:
        raise ValueError(f'About 시트의 필수 항목이 없습니다: {", ".join(sorted(missing))}')

    biography = data['biography']
    cv = data['cv']
    biography['paragraphs'] = []
    cv['researchInterests']['items'] = []
    cv['education']['items'] = []
    cv['peerReview']['journals'] = []
    cv['leadership']['blocks'] = []

    entries = workbook['About 목록'].iter_rows(values_only=True)
    if tuple(next(entries)[:7]) != (
        '구역', '기간', '제목·학위', '내용', '학교·설명', '이미지 경로', '이미지 설명'
    ):
        raise ValueError('About 목록 시트의 열 제목이 변경되었습니다.')

    count = 0
    for row_number, row in enumerate(entries, start=2):
        values = [as_text(cell) for cell in row[:7]]
        if not any(values):
            continue
        category, period, title, text, detail, image, image_alt = values
        if category not in CATEGORIES:
            raise ValueError(f'About 목록 {row_number}행: 구역을 선택하세요.')
        if category in {'소개 문단', '연구 분야', '심사 저널'} and not text:
            raise ValueError(f'About 목록 {row_number}행: 내용을 입력하세요.')
        if category == '소개 문단':
            biography['paragraphs'].append(text)
        elif category == '연구 분야':
            cv['researchInterests']['items'].append(text)
        elif category == '학력':
            if not (period and title and detail):
                raise ValueError(f'About 목록 {row_number}행: 학력의 기간·학위·학교를 입력하세요.')
            cv['education']['items'].append({
                'period': period, 'degree': title, 'schoolLine': detail,
                'logoSrc': image, 'logoAlt': image_alt,
            })
        elif category == '심사 저널':
            cv['peerReview']['journals'].append(text)
        elif category == '활동':
            if not (period and title and text):
                raise ValueError(f'About 목록 {row_number}행: 활동의 기간·제목·내용을 입력하세요.')
            blocks = cv['leadership']['blocks']
            if blocks and blocks[-1]['heading'] == title and blocks[-1]['period'] == period:
                blocks[-1]['bullets'].append(text)
            else:
                blocks.append({'heading': title, 'period': period, 'bullets': [text]})
        count += 1

    return data, count


def main():
    workbook = load_workbook(SOURCE, read_only=True, data_only=True)
    data, count = build_data(workbook)
    workbook.close()
    text = '// Generated from content/site-content.xlsx by scripts/build-about.py.\n'
    text += 'window.ABOUT_CONTENT = ' + json.dumps(data, ensure_ascii=False, indent=2) + ';\n'
    temporary = OUTPUT.with_name(OUTPUT.name + '.tmp')
    temporary.write_text(text, encoding='utf-8')
    temporary.replace(OUTPUT)
    print(f'Generated {OUTPUT.relative_to(ROOT)} from {count} list entries')


if __name__ == '__main__':
    main()
