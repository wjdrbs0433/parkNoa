"""Generate Abstract page data from content/site-content.xlsx."""

import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'content' / 'site-content.xlsx'
OUTPUT = ROOT / 'inc' / 'content' / 'abstract-content.js'
HEADERS = ('번호', '저자·연도', '초록 제목', '제목 뒤 내용')


def as_text(value):
    return '' if value is None else str(value).strip()


def build_data(workbook):
    if 'Abstract' not in workbook:
        raise ValueError('Abstract 시트가 없습니다.')
    rows = workbook['Abstract'].iter_rows(values_only=True)
    if tuple(next(rows)[:4]) != HEADERS:
        raise ValueError('Abstract 시트의 A~D열 제목이 변경되었습니다.')

    items = []
    seen = set()
    for row_number, row in enumerate(rows, start=2):
        values = list(row[:4]) + [None] * max(0, 4 - len(row))
        if not any(as_text(value) for value in values):
            continue
        number, before, title, after = values[:4]
        if isinstance(number, bool) or not isinstance(number, (int, float)) or int(number) != number or number < 1:
            raise ValueError(f'Abstract {row_number}행: A열 번호는 양의 정수여야 합니다.')
        number = int(number)
        if number in seen:
            raise ValueError(f'Abstract {row_number}행: 중복 번호 {number}')
        seen.add(number)
        before, title, after = map(as_text, (before, title, after))
        if not (before and title and after):
            raise ValueError(f'Abstract {row_number}행: B~D열을 모두 입력하세요.')
        items.append({'number': number, 'before': before, 'title': title, 'after': after})

    items.sort(key=lambda item: item['number'], reverse=True)
    return {'items': items}


def main():
    workbook = load_workbook(SOURCE, read_only=True, data_only=True)
    try:
        data = build_data(workbook)
    finally:
        workbook.close()
    text = '// Generated from content/site-content.xlsx by scripts/build-abstract.py.\n'
    text += 'window.ABSTRACT_CONTENT = ' + json.dumps(data, ensure_ascii=False, indent=2) + ';\n'
    temporary = OUTPUT.with_name(OUTPUT.name + '.tmp')
    temporary.write_text(text, encoding='utf-8')
    temporary.replace(OUTPUT)
    print(f'Generated {OUTPUT.relative_to(ROOT)} from {len(data["items"])} abstracts')


if __name__ == '__main__':
    main()
