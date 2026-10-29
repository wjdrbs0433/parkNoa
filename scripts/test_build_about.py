"""Checks that a new About 목록 row appears in generated page data."""

import unittest
import importlib.util
from pathlib import Path

from openpyxl import load_workbook

spec = importlib.util.spec_from_file_location('build_about', Path(__file__).with_name('build-about.py'))
build_about = importlib.util.module_from_spec(spec)
spec.loader.exec_module(build_about)
SOURCE = build_about.SOURCE
build_data = build_about.build_data


class SheetRows:
    def __init__(self, rows):
        self.rows = rows

    def iter_rows(self, values_only=True):
        return iter(self.rows)


class AboutWorkbookTest(unittest.TestCase):
    def test_new_research_interest_row_is_included(self):
        source = load_workbook(SOURCE, read_only=True, data_only=True)
        settings = list(source['About'].iter_rows(values_only=True))
        entries = list(source['About 목록'].iter_rows(values_only=True))
        source.close()

        entry = ('연구 분야', '', '', 'Additional research topic', '', '', '')
        entries.append(entry)
        workbook = {
            'About': SheetRows(settings),
            'About 목록': SheetRows(entries),
        }
        data, count = build_data(workbook)

        self.assertEqual(data['cv']['researchInterests']['items'][-1], entry[3])
        self.assertEqual(count, len(entries) - 1)

    def test_education_school_edit_is_included(self):
        source = load_workbook(SOURCE, read_only=True, data_only=True)
        settings = list(source['About'].iter_rows(values_only=True))
        entries = list(source['About 목록'].iter_rows(values_only=True))
        source.close()

        row_index = next(i for i, row in enumerate(entries) if row[1] == 'Aug 2026-Present')
        edited = list(entries[row_index])
        edited[4] = '1111' + edited[4]
        entries[row_index] = tuple(edited)
        workbook = {
            'About': SheetRows(settings),
            'About 목록': SheetRows(entries),
        }

        data, _ = build_data(workbook)
        self.assertEqual(data['cv']['education']['items'][0]['schoolLine'], edited[4])


if __name__ == '__main__':
    unittest.main()
